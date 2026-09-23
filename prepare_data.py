"""Rebuild aggregated dashboard data: python prepare_data.py path/to/workbook.xlsx"""
import argparse, collections, datetime, hashlib, json
from pathlib import Path
import openpyxl

def key(v):
    if v is None: return ''
    if isinstance(v, (int, float)) and float(v).is_integer(): return str(int(v))
    s=str(v).strip()
    return str(int(s)) if s.isdigit() else s

def generate(source, destination):
    w=openpyxl.load_workbook(source, read_only=True, data_only=True)
    maps={}
    for sheet in ['entidad','alcaldia']:
        maps[sheet]={}
        for row in w[sheet].iter_rows(min_row=2,values_only=True):
            if row[0] is None: continue
            code=key(row[0]);name=str(row[1]).strip()
            if code in maps[sheet]: raise ValueError(f'Duplicate lookup key: {sheet} {code}')
            maps[sheet][code]=name
    sheet=w['database'];it=sheet.iter_rows(values_only=True)
    headers=[str(v).strip().upper() for v in next(it)]
    positions=[headers.index(n) for n in ['CVE_MUN','SEXO','NACIONALIDAD','ENT_PAIS_NAC']]
    municipalities={};origins={};raw={k:collections.Counter() for k in ['sex','nat','birth']};total=0;blank=0;unmatched_m=set();unmatched_o=set()
    for row in it:
        if all(v is None for v in row): blank+=1;continue
        mun,sex,nat,origin=[key(row[i]) for i in positions]
        if mun not in maps['alcaldia']: unmatched_m.add(mun)
        if origin not in maps['entidad']: unmatched_o.add(origin)
        try: origin_number=int(origin)
        except ValueError: origin_number=-1
        group='cdmx' if origin=='9' else 'mexico' if 1<=origin_number<=32 else 'abroad' if 100<=origin_number<=536 and origin in maps['entidad'] else 'unknown'
        origins[origin]={'name':maps['entidad'].get(origin,f'Sin catálogo ({origin or "vacío"})'),'group':group}
        m=municipalities.setdefault(mun,{'code':mun,'name':maps['alcaldia'].get(mun,f'Sin catálogo ({mun or "vacío"})'),'n':0,'sex':[0,0,0],'nat':[0,0,0],'birth':[0,0,0,0],'origins':{}})
        si={'1':0,'3':1}.get(sex,2);ni={'1':0,'3':1}.get(nat,2);bi={'cdmx':0,'mexico':1,'abroad':2,'unknown':3}[group]
        m['n']+=1;m['sex'][si]+=1;m['nat'][ni]+=1;m['birth'][bi]+=1;m['origins'][origin]=m['origins'].get(origin,0)+1
        raw['sex'][sex]+=1;raw['nat'][nat]+=1;raw['birth'][origin]+=1;total+=1
    for m in municipalities.values():
        for dimension in ['sex','nat','birth']: assert sum(m[dimension])==m['n']
        assert sum(m['origins'].values())==m['n']
    assert sum(m['n'] for m in municipalities.values())==total
    data={'source':Path(source).name,'sourceSha256':hashlib.sha256(Path(source).read_bytes()).hexdigest(),'generated':datetime.datetime.now().strftime('%Y-%m-%d'),'total':total,'municipalities':list(municipalities.values()),'origins':origins,'audit':{'blankRows':blank,'unmatchedMunicipalities':len(unmatched_m),'unmatchedOrigins':len(unmatched_o),'rawCounts':{k:dict(v) for k,v in raw.items()}}}
    destination.mkdir(parents=True,exist_ok=True)
    (destination/'data.js').write_text('window.DASHBOARD_DATA = '+json.dumps(data,ensure_ascii=False,separators=(',',':'))+';\n',encoding='utf-8')
    (destination/'data-quality.json').write_text(json.dumps({k:v for k,v in data.items() if k not in ['municipalities','origins']},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'records':total,'alcaldias':len(municipalities),'birthplaceGroups':[sum(m['birth'][i] for m in municipalities.values()) for i in range(4)],'unmatchedMunicipalities':sorted(unmatched_m),'unmatchedOrigins':sorted(unmatched_o)}))

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('workbook',type=Path);parser.add_argument('--out',type=Path,default=Path(__file__).resolve().parent)
    args=parser.parse_args();generate(args.workbook,args.out)
    if args.out.resolve()==Path(__file__).resolve().parent:
        from build_dashboard import build
        build()
